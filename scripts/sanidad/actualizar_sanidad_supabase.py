# -*- coding: utf-8 -*-
"""
Lee los informes "Proximos"/"Ultimos" que la Planilla Sanitaria ya publica
como .xlsx (foto, no calcula nada nuevo) y los sube a dos tablas de Supabase
(sanidad_proximos / sanidad_ultimos) para que las apps de potreros los puedan
mostrar en un panel de solo lectura.

Reemplaza (delete + insert) las filas de cada establecimiento en cada corrida
-- no hace falta upsert incremental, es solo una foto mas.

Uso: python actualizar_sanidad_supabase.py
"""
import json
import os
import openpyxl
from datetime import datetime, date
from urllib import request, error

CARPETA = r"C:\Users\Pedro\OneDrive\A Registros PEDRO LANFRANCO CRESPO\Sanidad"
LOG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "actualizar_sanidad_supabase.log")
SUPA_URL = "https://skkknfjpwcstefcroqjt.supabase.co"
SUPA_KEY = "sb_publishable_U62qYc-mqby-ZGzE9FG2IA_kzOTGGJK"


def log(linea):
    with open(LOG, "a", encoding="utf-8") as f:
        f.write(f"{datetime.now().isoformat(timespec='seconds')} {linea}\n")

ARCHIVOS = [
    ("Proximos_celular.xlsx", "la_vuelta", "sanidad_proximos"),
    ("Ultimos_celular.xlsx", "la_vuelta", "sanidad_ultimos"),
    ("Proximos_ML_celular.xlsx", "maria_laura", "sanidad_proximos"),
    ("Ultimos_ML_celular.xlsx", "maria_laura", "sanidad_ultimos"),
]

# Catalogo de productos (P1.3 del plan de mejoras, 7/9/2026): a diferencia de
# lo de arriba, esto no viene de un .xlsx ya exportado -- se lee directo de
# la hoja "Costo productos" del .xlsm de La Vuelta (Maria Laura sincroniza su
# catalogo desde ahi, asi que uno solo alcanza). Alimenta los desplegables de
# Producto 1/2 en "+ Cargar tratamiento".
XLSM_LA_VUELTA = os.path.join(CARPETA, "Planilla_Sanitaria_v22.xlsm")
HOJA_CATALOGO = "Costo productos"


def a_iso(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return None


def leer_filas(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    encabezados = [str(h).strip() if h is not None else "" for h in filas[0]]
    out = []
    for fila in filas[1:]:
        if all(v is None for v in fila):
            continue
        out.append(dict(zip(encabezados, fila)))
    return out


def a_proximos(filas, establecimiento):
    filas_out = []
    for r in filas:
        filas_out.append({
            "establecimiento": establecimiento,
            "categoria": r.get("Categoría"),
            "sanidad": r.get("Sanidad"),
            "animales": r.get("Animales"),
            "ultimo_producto": r.get("Último producto"),
            "ultimo_tratamiento": a_iso(r.get("Últ. tratamiento")),
            "fecha_prevista": a_iso(r.get("Fecha prevista")),
            "prevista_txt": r.get("Prevista"),
            "dias": r.get("Días"),
            "estado": r.get("Estado"),
            "actualizado": a_iso(r.get("Actualizado")),
        })
    return filas_out


def a_ultimos(filas, establecimiento):
    filas_out = []
    for r in filas:
        filas_out.append({
            "establecimiento": establecimiento,
            "fecha": a_iso(r.get("Fecha")),
            "fecha_txt": r.get("Fecha txt"),
            "categoria": r.get("Categoría"),
            "potrero": str(r.get("Potrero")) if r.get("Potrero") is not None else None,
            "cantidad": r.get("Cantidad"),
            "productos": r.get("Productos"),
            "dosis": r.get("Dosis"),
            "estado": r.get("Estado"),
            "apto_desde": a_iso(r.get("Apto desde")),
            "apto_txt": r.get("Apto txt"),
            "dias_hace": r.get("Hace (días)"),
            "costo_animal": r.get("Costo x animal"),
            "observaciones": r.get("Observaciones"),
            "actualizado": a_iso(r.get("Actualizado")),
        })
    return filas_out


def _num(v):
    """Convierte a float si se puede; None si esta vacio o no es numero
    (deja el campo afuera del catalogo en vez de mandar un 0 enganoso)."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def leer_catalogo_productos(ruta):
    """Lee la hoja 'Costo productos' directo del .xlsm (read-only, no toca el
    archivo). Fila 4 = encabezados, datos desde la fila 5. Columnas fijas
    (B=Producto, C=Tipo, G=Categoria, L=T espera, N=Valor dosis $/ml) --
    confirmado leyendo la hoja a mano, no hay una fila de encabezados
    machine-readable mas arriba que valga la pena parsear para esto.

    L (T espera, indice 11) es el plazo para FAENA -- no confundir con K
    (T_Residual, indice 10), que es cada cuanto se REPITE el tratamiento.
    Son dos plazos distintos (ver memoria planilla-sanitaria-v22, Bug 11);
    esta funcion solo lee L a proposito.

    N (Valor dosis $/ml, indice 13) ya viene calculado por el Excel a partir
    de Importe/Presentacion -- se lee ese valor final tal cual, no se
    re-deriva la division acá (evita repetir el bug real de unidades
    ml/L que ya paso con Ectoraz/Tacplus)."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[HOJA_CATALOGO]
    productos = []
    vistos = set()
    for row in ws.iter_rows(min_row=5, values_only=True):
        nombre = row[1] if len(row) > 1 else None
        if not nombre:
            continue
        nombre = str(nombre).strip()
        if not nombre or nombre in vistos:
            continue
        vistos.add(nombre)
        dias_retiro = _num(row[11]) if len(row) > 11 else None
        productos.append({
            "producto": nombre,
            "tipo": (str(row[2]).strip() if len(row) > 2 and row[2] else None),
            "categoria": (str(row[6]).strip() if len(row) > 6 and row[6] else None),
            "dias_retiro": int(dias_retiro) if dias_retiro is not None else None,
            "valor_dosis_ml": _num(row[13]) if len(row) > 13 else None,
        })
    return productos


def rest(metodo, tabla, query="", body=None, extra_headers=None):
    url = f"{SUPA_URL}/rest/v1/{tabla}{query}"
    headers = {
        "apikey": SUPA_KEY,
        "Authorization": f"Bearer {SUPA_KEY}",
        "Content-Type": "application/json",
    }
    if extra_headers:
        headers.update(extra_headers)
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = request.Request(url, data=data, headers=headers, method=metodo)
    try:
        with request.urlopen(req) as resp:
            return resp.status, resp.read()
    except error.HTTPError as e:
        return e.code, e.read()


def _restaurar_respaldo(tabla, respaldo):
    """Best-effort: si un POST falla despues de haber borrado, intenta
    reinsertar lo que habia antes en vez de dejar la tabla vacia. No es
    perfecto (si esto tambien falla, ya se perdio) pero es mucho mejor que
    no intentarlo -- confirmado el 12/9/2026: un cambio de schema (columna
    nueva sin crear todavia) hizo fallar el POST de productos_catalogo
    DESPUES de que el DELETE ya habia corrido, y la tabla quedo vacia
    varias horas hasta notarlo."""
    if not respaldo:
        return
    status, cuerpo = rest("POST", tabla, "", body=respaldo,
                           extra_headers={"Prefer": "return=minimal"})
    if status not in (200, 201):
        log(f"ERROR CRITICO: no se pudo restaurar el respaldo de {tabla} ({len(respaldo)} filas) -- {status} {cuerpo}")
    else:
        log(f"restaurado el respaldo de {tabla} ({len(respaldo)} filas) tras un error")


def reemplazar(tabla, establecimiento, filas):
    status, cuerpo = rest("GET", tabla, f"?establecimiento=eq.{establecimiento}")
    respaldo = json.loads(cuerpo) if status == 200 else []

    status, cuerpo = rest("DELETE", tabla, f"?establecimiento=eq.{establecimiento}",
                           extra_headers={"Prefer": "return=minimal"})
    if status not in (200, 204):
        raise RuntimeError(f"Error borrando {tabla}/{establecimiento}: {status} {cuerpo}")
    if not filas:
        return 0
    status, cuerpo = rest("POST", tabla, "", body=filas,
                           extra_headers={"Prefer": "return=minimal"})
    if status not in (200, 201):
        _restaurar_respaldo(tabla, respaldo)
        raise RuntimeError(f"Error insertando en {tabla}/{establecimiento}: {status} {cuerpo}")
    return len(filas)


def reemplazar_catalogo(productos):
    status, cuerpo = rest("GET", "productos_catalogo", "?select=*")
    respaldo = json.loads(cuerpo) if status == 200 else []

    status, cuerpo = rest("DELETE", "productos_catalogo", "?producto=not.is.null",
                           extra_headers={"Prefer": "return=minimal"})
    if status not in (200, 204):
        raise RuntimeError(f"Error borrando productos_catalogo: {status} {cuerpo}")
    if not productos:
        return 0
    status, cuerpo = rest("POST", "productos_catalogo", "", body=productos,
                           extra_headers={"Prefer": "return=minimal"})
    if status not in (200, 201):
        _restaurar_respaldo("productos_catalogo", respaldo)
        raise RuntimeError(f"Error insertando en productos_catalogo: {status} {cuerpo}")
    return len(productos)


if __name__ == "__main__":
    log("--- corrida iniciada ---")
    for nombre_archivo, establecimiento, tabla in ARCHIVOS:
        ruta = f"{CARPETA}\\{nombre_archivo}"
        try:
            filas_crudas = leer_filas(ruta)
        except FileNotFoundError:
            log(f"(no existe {nombre_archivo}, se salteo)")
            continue
        except Exception as e:
            log(f"ERROR leyendo {nombre_archivo}: {e}")
            continue
        if tabla == "sanidad_proximos":
            filas = a_proximos(filas_crudas, establecimiento)
        else:
            filas = a_ultimos(filas_crudas, establecimiento)
        try:
            n = reemplazar(tabla, establecimiento, filas)
            log(f"{tabla} / {establecimiento}: {n} filas (desde {nombre_archivo})")
        except Exception as e:
            log(f"ERROR subiendo {tabla}/{establecimiento}: {e}")

    try:
        productos = leer_catalogo_productos(XLSM_LA_VUELTA)
        n = reemplazar_catalogo(productos)
        log(f"productos_catalogo: {n} productos (desde {os.path.basename(XLSM_LA_VUELTA)})")
    except FileNotFoundError:
        log(f"(no existe {XLSM_LA_VUELTA}, se salteo el catalogo)")
    except Exception as e:
        log(f"ERROR subiendo productos_catalogo: {e}")
